import tkinter as tk
from tkinter import messagebox, filedialog
from docx import Document
from docx.shared import Pt

class WordCloneApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Word Clone - Éditeur de Structure")
        self.doc = None
        self.file_path = None

        # --- PANNEAU SUPÉRIEUR ---
        top_frame = tk.Frame(root)
        top_frame.pack(pady=10, fill=tk.X)
        tk.Button(top_frame, text="📁 Ouvrir Word", command=self.load_document).pack(side=tk.LEFT, padx=10)
        tk.Button(top_frame, text="💾 Sauvegarder", command=self.save_doc, bg="#28a745", fg="white").pack(side=tk.LEFT)

        # --- LISTE ET BOUTONS DE DÉPLACEMENT ---
        list_container = tk.Frame(root)
        list_container.pack(pady=5, padx=10)

        self.listbox = tk.Listbox(list_container, width=85, height=12, selectmode=tk.SINGLE)
        self.listbox.pack(side=tk.LEFT)
        self.listbox.bind('<<ListboxSelect>>', self.on_select_line)

        move_frame = tk.Frame(list_container)
        move_frame.pack(side=tk.LEFT, padx=5)
        tk.Button(move_frame, text="▲", command=lambda: self.move_paragraph(-1)).pack(pady=2)
        tk.Button(move_frame, text="▼", command=lambda: self.move_paragraph(1)).pack(pady=2)

        # --- ZONE D'ÉDITION ---
        edit_frame = tk.Frame(root)
        edit_frame.pack(pady=10)
        self.entry_edit = tk.Entry(edit_frame, width=70)
        self.entry_edit.pack(pady=5)

        # --- BOUTONS D'ACTION ---
        action_frame = tk.Frame(root)
        action_frame.pack(pady=5)
        tk.Button(action_frame, text="Modifier", command=self.apply_text_change, width=12).grid(row=0, column=0, padx=5)
        tk.Button(action_frame, text="Insérer après", command=self.insert_paragraph_after, bg="#0078d4", fg="white", width=12).grid(row=0, column=1, padx=5)
        tk.Button(action_frame, text="Supprimer", command=self.delete_paragraph, bg="#dc3545", fg="white", width=12).grid(row=0, column=2, padx=5)

        # --- FORMATAGE ---
        fmt_frame = tk.Frame(root)
        fmt_frame.pack(pady=10)
        tk.Button(fmt_frame, text="Gras", command=lambda: self.apply_style('bold'), width=8).grid(row=0, column=0, padx=2)
        tk.Button(fmt_frame, text="Italique", command=lambda: self.apply_style('italic'), width=8).grid(row=0, column=1, padx=2)

    def load_document(self):
        self.file_path = filedialog.askopenfilename(filetypes=[("Word files", "*.docx")])
        if self.file_path:
            self.doc = Document(self.file_path)
            self.refresh_listbox()

    def refresh_listbox(self, select_index=None):
        self.listbox.delete(0, tk.END)
        for para in self.doc.paragraphs:
            txt = para.text if para.text.strip() else "[Ligne vide / Objet]"
            self.listbox.insert(tk.END, txt)
        if select_index is not None:
            self.listbox.select_set(select_index)
            self.listbox.see(select_index)

    def on_select_line(self, event):
        selection = self.listbox.curselection()
        if selection:
            index = selection[0]
            self.entry_edit.delete(0, tk.END)
            self.entry_edit.insert(0, self.doc.paragraphs[index].text)

    def apply_text_change(self):
        selection = self.listbox.curselection()
        if selection:
            index = selection[0]
            self.doc.paragraphs[index].text = self.entry_edit.get()
            self.refresh_listbox(index)

    def insert_paragraph_after(self):
        selection = self.listbox.curselection()
        if not selection: return
        index = selection[0]
        new_para = self.doc.paragraphs[index].insert_paragraph_after(self.entry_edit.get())
        new_para.style = self.doc.paragraphs[index].style
        self.refresh_listbox(index + 1)

    def delete_paragraph(self):
        selection = self.listbox.curselection()
        if not selection: return
        index = selection[0]
        p = self.doc.paragraphs[index]._element
        p.getparent().remove(p)
        self.refresh_listbox()

    def move_paragraph(self, direction):
        selection = self.listbox.curselection()
        if not selection: return
        old_index = selection[0]
        new_index = old_index + direction

        if 0 <= new_index < len(self.doc.paragraphs):
            p_element = self.doc.paragraphs[old_index]._element
            if direction == -1: # Monter
                p_element.getparent().insert(new_index, p_element)
            else: # Descendre
                # Pour descendre, on insère après la cible
                target_element = self.doc.paragraphs[new_index]._element
                target_element.addnext(p_element)
            
            self.refresh_listbox(new_index)

    def apply_style(self, mode):
        selection = self.listbox.curselection()
        if selection:
            index = selection[0]
            for run in self.doc.paragraphs[index].runs:
                if mode == 'bold': run.bold = True
                elif mode == 'italic': run.italic = True
            messagebox.showinfo("Style", "Appliqué !")

    def save_doc(self):
        if self.doc:
            path = filedialog.asksaveasfilename(defaultextension=".docx")
            if path:
                self.doc.save(path)
                messagebox.showinfo("Succès", "Sauvegardé !")

if __name__ == "__main__":
    root = tk.Tk()
    app = WordCloneApp(root)
    root.mainloop()
